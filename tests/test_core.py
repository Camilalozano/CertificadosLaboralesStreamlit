import csv
import io
import unittest
from datetime import date
from unittest.mock import patch
from docx import Document
from src.contracts import parse_csv, search, REFERENCE, OBJECT
from src.certificate import build_certificate, formatear_valor_pesos, valor_documento_texto
from src.documents import list_documents, ranked_pdfs
from src.network import SourceError, download_oracle, download_pdf, SameHostRedirect
from src.obligations import extract
from src.workflow import audit
from urllib.request import Request


def csv_fixture():
    stream = io.StringIO(newline='')
    writer = csv.writer(stream)
    writer.writerow([REFERENCE, OBJECT, 'id_contrato', 'proveedor_adjudicado', 'documento_proveedor',
                     'valor_del_contrato', 'fecha_de_inicio_del_contrato', 'fecha_de_fin_del_contrato'])
    writer.writerow(['ATENEA-003-2025', 'Análisis de información y gestión', 'CO1.PCCNTR.100', 'Persona de prueba', '001234', '90000000.0', '2025-01-10', '2025-10-10'])
    writer.writerow(['ATENEA-003-2026', 'Apoyo jurídico contractual', 'CO1.PCCNTR.200', 'Otra persona', '005678', '50000000', '2026-01-10', '2026-10-10'])
    writer.writerow(['ATENEA-312-2025', 'Análisis de información', 'CO1.PCCNTR.300', 'Persona tercera', '000987', '40000000', '', ''])
    return stream.getvalue().encode('utf-8-sig')


class ContractTests(unittest.TestCase):
    def setUp(self):
        self.catalog = parse_csv(csv_fixture())

    def test_exact_contract_and_id(self):
        self.assertEqual(search(self.catalog, 'atenea 003 2025', 'Número del contrato')[0].fields['id_contrato'], 'CO1.PCCNTR.100')
        self.assertEqual(len(search(self.catalog, 'CO1.PCCNTR.100', 'Número del contrato')), 1)

    def test_number_ambiguity_and_segment(self):
        self.assertEqual(len(search(self.catalog, '3', 'Número del contrato')), 2)
        self.assertEqual(len(search(self.catalog, '12', 'Número del contrato')), 0)

    def test_object_accents_and_tokens(self):
        self.assertEqual(len(search(self.catalog, 'informacion analisis', 'Objeto del contrato')), 2)
        self.assertEqual(search(self.catalog, '', 'Objeto del contrato'), [])

    def test_missing_or_duplicate_headers(self):
        for data in [b'a,b\n1,2', f'{REFERENCE},{REFERENCE},{OBJECT}\nx,x,x'.encode()]:
            with self.assertRaises(SourceError): parse_csv(data)

    def test_identification_preserves_leading_zeroes(self):
        self.assertEqual(self.catalog.contracts[0].fields['documento'], '001234')
        self.assertEqual(valor_documento_texto('001234.0'), '001234')

    def test_duplicate_rows_are_not_silently_selected(self):
        data = csv_fixture()
        data += data.decode('utf-8-sig').splitlines()[1].encode() + b'\n'
        self.assertEqual(len(search(parse_csv(data), 'ATENEA-003-2025', 'Número del contrato')), 2)

    def test_bare_duration_does_not_invent_days_or_months(self):
        raw = f'{REFERENCE},{OBJECT},duraci_n_del_contrato (contratos_electronicos)\nATENEA-1-2025,Objeto de prueba,243\n'
        self.assertEqual(parse_csv(raw.encode()).contracts[0].fields['plazo'], '243 (unidad por confirmar)')


class CertificateTests(unittest.TestCase):
    def setUp(self): self.fields = parse_csv(csv_fixture()).contracts[0].fields

    def test_money_from_oracle_and_localized(self):
        for value in ['90000000.0', '90000000', '90.000.000', '90,000,000.00']:
            self.assertEqual(formatear_valor_pesos(value), '$90.000.000 M/CTE')
        self.assertEqual(formatear_valor_pesos('1200.50'), '$1.200,50 M/CTE')
        self.assertEqual(formatear_valor_pesos('0'), '$0 M/CTE')
        for invalid in ['NaN', '-1', 'un millón']:
            with self.assertRaises(SourceError): formatear_valor_pesos(invalid)

    def test_docx_contains_correct_fields_obligations_and_unsigned_status(self):
        doc = Document(io.BytesIO(build_certificate(self.fields, ['Apoyar el análisis.'], date(2026, 10, 2))))
        content = '\n'.join(p.text for p in doc.paragraphs)
        self.assertIn('001234', content)
        self.assertIn('$90.000.000', content)
        self.assertIn('1. Apoyar el análisis.', content)
        self.assertIn('Firma del funcionario autorizado', content)
        self.assertIn('BORRADOR', doc.sections[0].header.paragraphs[0].text)
        self.assertNotIn('CAMILO CARDOZO CRUZ', content)

    def test_required_fields_and_date_order(self):
        with self.assertRaises(SourceError): build_certificate(self.fields, [], date.today())
        with self.assertRaises(SourceError): build_certificate({**self.fields, 'fin': '2024-01-01'}, ['Texto'], date.today())
        with self.assertRaises(SourceError): build_certificate({**self.fields, 'documento': ''}, ['Texto'], date.today())

    def test_audit_marks_manual_edits_and_excludes_secret(self):
        c = parse_csv(csv_fixture()).contracts[0]
        result = audit(c, 'digest', {**c.fields, 'contratista': 'Corrección'}, ['Texto cambiado'],
                       [{'texto': 'Original', 'pagina': 2}], 'minuta.pdf', 'pdfhash', '', date.today(), '', '')
        self.assertEqual(result['campos_editados'], ['contratista'])
        self.assertEqual(result['obligaciones_certificado'][0]['origen'], 'Edición manual')
        self.assertNotIn('PAR', str(result))


class DocumentTests(unittest.TestCase):
    @patch('src.documents.api_rows')
    def test_pagination_and_contract_filter(self, api):
        api.side_effect = [[{'id_documento': str(i), 'n_mero_de_contrato': 'CO1.PCCNTR.100'} for i in range(500)],
                           [{'id_documento': '500', 'n_mero_de_contrato': 'CO1.PCCNTR.100'}]]
        self.assertEqual(len(list_documents('CO1.PCCNTR.100')), 501)
        self.assertEqual(api.call_args.args[0]['$offset'], 500)

    @patch('src.documents.api_rows', return_value=[{'n_mero_de_contrato': 'CO1.PCCNTR.200'}])
    def test_wrong_contract_rejected(self, api):
        with self.assertRaises(SourceError): list_documents('CO1.PCCNTR.100')

    def test_ranking_does_not_prefer_accounts_or_modifications(self):
        rows = [{'nombre_archivo': n, 'extensi_n': 'pdf', 'id_documento': str(i)} for i, n in enumerate([
            'CUENTA ATENEA-003-2025.pdf', 'MINUTA ATENEA-003-2025.pdf', 'MODIFICACION MINUTA ATENEA-003-2025.pdf'])]
        self.assertEqual(ranked_pdfs(rows, 'ATENEA-003-2025')[0][1]['id_documento'], '1')

    def test_network_destination_limits(self):
        for url in ['http://localhost/data', 'https://objectstorage.us-ashburn-1.oraclecloud.com.evil.test/p/x']:
            with self.assertRaises(SourceError): download_oracle(url)
        with self.assertRaises(SourceError): download_pdf('https://example.com/minuta.pdf')
        with self.assertRaises(SourceError):
            SameHostRedirect().redirect_request(Request('https://community.secop.gov.co/a'), None, 302, '', {}, 'http://localhost/file')


class ExtractionTests(unittest.TestCase):
    @patch('src.obligations.pdf_pages')
    def test_specific_only_and_page_provenance(self, pages):
        pages.return_value = ['OBLIGACIONES GENERALES: 1. Regla general.\nB) OBLIGACIONES ESPECÍFICAS: 1. Apoyar análisis.\n',
                              '2. Entregar informes.\nCLÁUSULA TERCERA: OBLIGACIONES DE LA CONTRATANTE: 1. Pagar.']
        result = extract(b'fixture')
        self.assertEqual([o['texto'] for o in result.obligations], ['Apoyar análisis.', 'Entregar informes.'])
        self.assertEqual([o['pagina'] for o in result.obligations], [1, 2])

    @patch('src.obligations.pdf_pages', return_value=['OBLIGACIONES GENERALES: 1. Pagar.'])
    def test_no_specific_section_does_not_invent(self, pages):
        with self.assertRaises(SourceError): extract(b'fixture')

    @patch('src.obligations.pdf_pages', return_value=['OBLIGACIONES ESPECIFICAS: 1. Apoyar. 3. Entregar. CLÁUSULA TERCERA: Fin.'])
    def test_numbering_gap_is_flagged(self, pages):
        self.assertTrue(extract(b'fixture').warnings)


if __name__ == '__main__': unittest.main()
