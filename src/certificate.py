"""Generación Word adaptada de generador_masivo_certificadoslaborales.py.

Conserva la estructura de ATENEA y corrige moneda, identificación y salidas en memoria.
"""
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from .network import SourceError, allowed_url, DOCUMENT_HOST

MONTHS = ('enero febrero marzo abril mayo junio julio agosto septiembre octubre noviembre diciembre').split()


def valor_documento_texto(value):
    value = str(value or '').strip()
    return re.sub(r'^(\d+)\.0+$', r'\1', value)


def formatear_valor_pesos(value):
    raw = str(value or '').strip()
    if not raw:
        return 'No informado'
    # Oracle usa punto decimal. Solo los valores con grupos de tres se interpretan como miles.
    raw = re.sub(r'(?i)M/CTE\.?|COP|PESOS|\$|\s', '', raw)
    if re.fullmatch(r'-?\d{1,3}(?:\.\d{3})+(?:,\d{1,2})?', raw):
        raw = raw.replace('.', '').replace(',', '.')
    elif re.fullmatch(r'-?\d{1,3}(?:,\d{3})+(?:\.\d{1,2})?', raw):
        raw = raw.replace(',', '')
    elif ',' in raw and '.' not in raw:
        raw = raw.replace(',', '.')
    try:
        amount = Decimal(raw)
        if not amount.is_finite() or amount < 0:
            raise InvalidOperation
        amount = amount.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    except InvalidOperation:
        raise SourceError('El valor del contrato no es un número válido. Revíselo antes de generar.') from None
    parts = f'{amount:,.2f}'.split('.')
    result = parts[0].replace(',', '.')
    if parts[1] != '00':
        result += ',' + parts[1]
    return '$' + result + ' M/CTE'


def parse_date(value):
    value = str(value or '').split('((')[0].strip()
    if not value:
        return None
    try:
        return datetime.fromisoformat(value.replace('Z', '+00:00')).date()
    except ValueError:
        for fmt in ('%d/%m/%Y', '%Y/%m/%d', '%d-%m-%Y'):
            try:
                return datetime.strptime(value, fmt).date()
            except ValueError:
                pass
    raise SourceError('Use fechas AAAA-MM-DD o DD/MM/AAAA, o deje el campo vacío si no está informado.')


def formatear_fecha_larga(value):
    day = parse_date(value)
    return f'{day.day} de {MONTHS[day.month-1]} de {day.year}' if day else 'No informada'


def limpiar_nombre_archivo(value):
    value = re.sub(r'[\\/*?:"<>|\x00-\x1f]', '_', str(value))
    return re.sub(r'\s+', '_', value).strip('._ ')[:130] or 'contrato'


def agregar_parrafo(doc, value='', bold=False, size=11, align=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(value)
    run.bold = bold
    run.font.name = 'Arial'
    run.font.size = Pt(size)
    return paragraph


def agregar_campo(doc, label, value):
    p = agregar_parrafo(doc)
    p.add_run(label + ': ').bold = True
    p.add_run(str(value or 'No informado'))


def hyperlink(paragraph, label, url):
    if not allowed_url(url, DOCUMENT_HOST):
        return
    relation = paragraph.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), relation)
    run = OxmlElement('w:r'); prop = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '00549F'); prop.append(color)
    run.append(prop); t = OxmlElement('w:t'); t.text = label; run.append(t); h.append(run); paragraph._p.append(h)


def build_certificate(fields, obligations, issued_on: date, signer='', role='', source_name='', source_url='', logo=None):
    for key, label in [('referencia', 'número del contrato'), ('contratista', 'nombre del contratista'),
                       ('documento', 'identificación'), ('objeto', 'objeto')]:
        if not str(fields.get(key, '')).strip():
            raise SourceError('Complete el campo ' + label + '.')
    if not obligations or not all(str(o).strip() for o in obligations):
        raise SourceError('Incluya las obligaciones específicas antes de generar el certificado.')
    first, last = parse_date(fields.get('inicio')), parse_date(fields.get('fin'))
    if first and last and last < first:
        raise SourceError('La fecha de finalización es anterior a la de inicio.')
    money = formatear_valor_pesos(fields.get('valor'))
    doc = Document()
    normal = doc.styles['Normal']; normal.font.name = 'Arial'; normal.font.size = Pt(11)
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.8)
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    header = section.header.paragraphs[0]
    header.add_run('BORRADOR PARA REVISIÓN Y FIRMA').font.size = Pt(8)
    if logo:
        try:
            doc.add_picture(io.BytesIO(logo), width=Inches(1.45))
        except Exception:
            raise SourceError('No fue posible usar el logo. Cargue una imagen PNG o JPG válida.') from None
    agregar_parrafo(doc, 'LA SUBGERENCIA DE GESTIÓN ADMINISTRATIVA DE LA AGENCIA DISTRITAL PARA LA\n'
                     'EDUCACIÓN SUPERIOR, LA CIENCIA Y LA TECNOLOGÍA - ATENEA', True, align=WD_ALIGN_PARAGRAPH.CENTER)
    agregar_parrafo(doc, 'CERTIFICA QUE:', True, align=WD_ALIGN_PARAGRAPH.CENTER)
    tipo = str(fields.get('tipo_documento') or 'documento de identificación').strip()
    intro = (f'Revisados los registros de contratación, se encontró que {fields["contratista"].upper()}, '
             f'con identificación {tipo} No. {valor_documento_texto(fields["documento"])}, '
             'suscribió con la Agencia Distrital para la Educación Superior, la Ciencia y la Tecnología '
             '- ATENEA, en calidad de contratista, el siguiente contrato:')
    agregar_parrafo(doc, intro, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    agregar_parrafo(doc, 'CONTRATO No. ' + fields['referencia'], True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for label, value in [('Objeto', fields['objeto']), ('Valor total reportado', money),
                         ('Fecha de inicio reportada', formatear_fecha_larga(fields.get('inicio'))),
                         ('Fecha de finalización reportada', formatear_fecha_larga(fields.get('fin'))),
                         ('Plazo contractual reportado', fields.get('plazo')),
                         ('Estado reportado', fields.get('estado'))]:
        agregar_campo(doc, label, value)
    agregar_parrafo(doc, 'Obligaciones específicas del contratista:', True)
    for index, obligation in enumerate(obligations, 1):
        p = agregar_parrafo(doc, f'{index}. {obligation.strip()}', size=10.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        p.paragraph_format.space_after = Pt(4)
    if source_name:
        agregar_parrafo(doc, 'Fuente de las obligaciones: ' + source_name, size=9)
    if source_url:
        hyperlink(agregar_parrafo(doc, size=9), 'Consultar documento de origen en SECOP II', source_url)
    agregar_parrafo(doc, f'Bogotá D. C., {issued_on.day} de {MONTHS[issued_on.month-1]} de {issued_on.year}.')
    agregar_parrafo(doc, 'Cordialmente,')
    agregar_parrafo(doc, '\n')
    agregar_parrafo(doc, signer.strip().upper() or 'Firma del funcionario autorizado', True)
    agregar_parrafo(doc, role.strip() or 'Subgerencia de Gestión Administrativa')
    agregar_parrafo(doc, 'Proyectó: ____________________    Revisó: ____________________', size=8)
    p = agregar_parrafo(doc, 'Documento generado mediante consulta de la base contractual y extracción automatizada. '
                       'Su contenido requiere revisión contra las fuentes oficiales y firma del funcionario competente.', size=9)
    for r in p.runs:
        r.font.color.rgb = RGBColor.from_string('555555')
    hyperlink(agregar_parrafo(doc, size=9), 'Consultar proceso en SECOP II', fields.get('url', ''))
    footer = section.footer.paragraphs[0]
    footer.text = 'Cr 10 # 28-49. Torre A, piso 26.\nBogotá D. C. Colombia\n(601) 666 0006\nwww.agenciaatenea.gov.co'
    for run in footer.runs:
        run.font.name = 'Arial'; run.font.size = Pt(8)
    output = io.BytesIO(); doc.save(output)
    return output.getvalue()
